#!/usr/bin/env python3
"""Build CPS 240 Model Risk Management — the draft cross-industry prudential standard.

Bold paragraphs are mandatory requirements. Everything else is explanatory.
Cross-references into CPG 240 are computed from the guide's own paragraph
numbering, so the two documents cannot fall out of step.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).parent))

import content_guide as G
import content_requirements as R
import content_sources as S
from docx_common import (
    AMBER, GREY, NAVY, SLATE, audit_fonts, banner, bottom_rule, bullets,
    callout, force_arial, landscape_section, page_break, para, portrait_section,
    page_setup, add_header_footer, build_styles, table, xref_box,
)

OUT = Path(__file__).resolve().parents[1] / "CPS_240_Model_Risk_Management.docx"
AS_OF = "25 July 2026"
STATUS = "POLICY DEVELOPMENT DRAFT — NOT IN FORCE"


def cover(doc):
    p = para(doc, "AUSTRALIAN PRUDENTIAL REGULATION AUTHORITY", "DocMeta")
    p.runs[0].font.size = Pt(9)
    p.runs[0].bold = True

    p = para(doc, "Prudential Standard CPS 240", "DocTitle")
    bottom_rule(p)
    para(doc, "Model Risk Management", "DocSubtitle").runs[0].bold = True
    para(doc, "Draft cross-industry prudential standard", "DocMeta")
    para(doc, f"Policy development version — information verified to {AS_OF}", "DocMeta")
    doc.add_paragraph(style="DocMeta")

    banner(doc, STATUS)

    callout(
        doc,
        "Status and legal review",
        "This document is a policy-development draft and is not in force. It has no legal effect. "
        "APRA Legal and the Office of Parliamentary Counsel must settle the enabling provisions, "
        "sectoral application, instrument numbering, definitions, commencement, transition, "
        "enforcement powers and interaction with existing CPS and SPS instruments before "
        "consultation or issue. Annex F records the decisions required. The number 'CPS 240' is a "
        "placeholder adopted for drafting convenience and does not indicate an allocated instrument "
        "number.",
    )


def objectives(doc):
    doc.add_paragraph("Objectives and key requirements", style="Heading 1")
    para(doc,
         "The objective of this Prudential Standard is to ensure that an APRA-regulated entity "
         "identifies the models it relies upon, manages the risk that those models are wrong or are "
         "used incorrectly, and can demonstrate that reliance on models is consistent with its "
         "prudential obligations.",
         "Guidance")
    para(doc,
         "Models determine regulatory capital, provisions, valuations, insurance liabilities, unit "
         "prices, credit and pricing decisions, fraud interventions and member outcomes. Model risk "
         "is the risk of adverse consequences from decisions based on incorrect or misused model "
         "outputs. It is a distinct prudential risk and is not adequately managed as a subset of the "
         "risks the models themselves measure.",
         "Guidance")
    para(doc,
         "This Prudential Standard is intended to operate alongside CPS 220 or SPS 220, CPS 230 and "
         "CPS 234, and alongside the model requirements in the capital, actuarial and investment "
         "governance standards. It does not displace those requirements, nor obligations under other "
         "Commonwealth, State or Territory laws or the responsibilities of other regulators.",
         "Guidance")
    para(doc,
         "The key requirements of this Prudential Standard are that an APRA-regulated entity must:",
         "Guidance")
    bullets(doc, [
        "identify the models it uses and maintain an accurate enterprise inventory;",
        "classify models by risk and apply controls of an intensity that matches that classification;",
        "govern models across their full lifecycle, from development through to decommissioning;",
        "subject models to independent validation and to effective challenge;",
        "control how models are used, changed and adjusted;",
        "monitor model performance and act on deterioration;",
        "manage model risk arising from third parties and from artificial intelligence techniques;",
        "assess model risk in aggregate and report it to the Board; and",
        "maintain records and obtain independent assurance sufficient to demonstrate compliance.",
    ])


def interpretation(doc):
    doc.add_paragraph("Interpretation", style="Heading 1")
    callout(
        doc,
        "How to read this Prudential Standard",
        "Every numbered paragraph set in bold is a proposed mandatory requirement. Text that is not "
        "bold is explanatory and does not itself impose an obligation. Each requirement is followed "
        "by explanatory text and by a cross-reference to the paragraphs of CPG 240 Model Risk "
        "Management that provide guidance on meeting it.",
    )
    para(doc,
         "Proportionality. References in this Prudential Standard to materiality, risk, complexity "
         "and criticality are cumulative considerations, applied to the circumstances of the entity "
         "and of the individual model.",
         "Guidance")
    para(doc,
         "Interaction with other standards. Where a model is also subject to a requirement in a "
         "capital, actuarial or investment governance standard, that requirement continues to apply. "
         "Where a model failure is also an operational risk incident or an information security "
         "incident, the notification requirements of CPS 230 and CPS 234 apply and this Prudential "
         "Standard does not create a separate deadline.",
         "Guidance")
    para(doc,
         "Terms defined in Annex A have the meaning given in that Annex.",
         "Guidance")


def contents(doc):
    doc.add_paragraph("Contents", style="Heading 1")
    seen = []
    for r in R.REQUIREMENTS:
        if r["part"] not in seen:
            seen.append(r["part"])
    for part in seen:
        p = para(doc, part, "Guidance")
        p.runs[0].bold = True
        secs = []
        for r in R.REQUIREMENTS:
            if r["part"] == part and r["section"] not in secs:
                secs.append(r["section"])
        for s in secs:
            ids = [r["id"] for r in R.REQUIREMENTS if r["section"] == s]
            para(doc, f"    {s}  ({ids[0]}–{ids[-1]})" if len(ids) > 1
                 else f"    {s}  ({ids[0]})", "Guidance")
    doc.add_paragraph(style="DocMeta")
    for a in ["Annex A — Definitions",
              "Annex B — Minimum model inventory fields",
              "Annex C — Model risk tiering indicators",
              "Annex D — International comparison of model risk requirements",
              "Annex E — Source register",
              "Annex F — Matters requiring legal and policy settlement"]:
        p = para(doc, a, "Guidance")
        p.runs[0].bold = True


def body(doc, req_map):
    current_part = None
    current_section = None

    for r in R.REQUIREMENTS:
        if r["part"] != current_part:
            page_break(doc)
            doc.add_paragraph(r["part"], style="Heading 1")
            current_part = r["part"]
            current_section = None

        if r["section"] != current_section:
            doc.add_paragraph(r["section"], style="Heading 2")
            current_section = r["section"]

        doc.add_paragraph(f"{r['id']} — {r['title']}", style="Heading 3")

        p = doc.add_paragraph(style="Requirement")
        run = p.add_run(f"{r['id']}. {r['requirement']}")
        run.bold = True

        for g in r["guidance"]:
            para(doc, g, "Guidance")

        nums = req_map.get(r["id"], [])
        if nums:
            xr = para(doc, f"→ Guidance: CPG 240 paragraphs {G.ranges(nums)}.", "CrossRef")
            xr.runs[0].italic = True


def annex_a(doc):
    page_break(doc)
    doc.add_paragraph("Annex A — Definitions", style="Heading 1")
    para(doc,
         "This Annex is intended to form part of the final Prudential Standard. Definitions are "
         "subject to legal drafting and to alignment with Australian Government and APRA "
         "terminology.",
         "Guidance")
    table(doc, ["Term", "Draft definition"], S.DEFINITIONS, [4.0, 12.6])


def annex_b(doc):
    page_break(doc)
    doc.add_paragraph("Annex B — Minimum model inventory fields", style="Heading 1")
    para(doc,
         "An APRA-regulated entity's model inventory must record at least the following for each "
         "model. These are minimum fields; an entity should record whatever further information its "
         "framework requires.",
         "Guidance")
    table(doc, ["Field group", "Minimum content"], S.INVENTORY_FIELDS, [4.0, 12.6])


def annex_c(doc):
    page_break(doc)
    doc.add_paragraph("Annex C — Model risk tiering indicators", style="Heading 1")
    para(doc,
         "An entity would ordinarily assign a model to its highest risk tier where one or more of "
         "the following indicators is material. Classification remains outcome-based and "
         "context-specific, and no single indicator is determinative.",
         "Guidance")
    for b in S.TIER_INDICATORS:
        para(doc, "• " + b, "Bullet")


def annex_d(doc):
    """The wide international comparison. Landscape so the columns are readable."""
    landscape_section(doc)
    doc.add_paragraph("Annex D — International comparison of model risk requirements",
                      style="Heading 1")
    para(doc,
         "This Annex compares the model risk requirements and expectations of the authorities whose "
         "frameworks informed this Prudential Standard. Legal status differs materially between "
         "them and is stated in each cell group. An overseas instrument is not binding in Australia "
         "whatever its status in its home jurisdiction.",
         "Guidance")
    para(doc,
         "Status codes: [B] binding in the home jurisdiction · [SE] supervisory expectation, not "
         "binding law · [G] guidance · [PB] proposed or consultative · [A] analytical or advisory · "
         "[—] not addressed.",
         "SourceNote")

    headers = ["Domain", "APRA (proposed CPS 240)", "OCC / Federal Reserve (US)",
               "PRA (UK)", "OSFI (Canada)", "ECB / SSM (EU)", "MAS (Singapore)",
               "BCBS", "FSB"]
    widths = [2.5, 3.1, 3.0, 3.0, 3.0, 3.0, 2.9, 2.7, 2.5]
    table(doc, headers, S.CROSSWALK_ROWS, widths, font_pt=7)

    portrait_section(doc)


def annex_e(doc):
    doc.add_paragraph("Annex E — Source register", style="Heading 1")
    para(doc,
         "Every source relied upon in this Prudential Standard and in CPG 240. URLs are to official "
         "authority websites and were checked on the date shown. Where a source could not be "
         "confirmed from an official domain it is marked accordingly and has not been relied upon "
         "for any statement of requirement.",
         "Guidance")
    rows = [[s["id"], s["authority"], s["title"], s["date"], s["status"], s["url"]]
            for s in S.SOURCES]
    table(doc, ["ID", "Authority", "Title", "Date / effect", "Legal status", "Official URL"],
          rows, [1.9, 1.5, 4.4, 2.3, 3.0, 3.5], font_pt=7)


def annex_f(doc):
    page_break(doc)
    doc.add_paragraph("Annex F — Matters requiring legal and policy settlement", style="Heading 1")
    para(doc,
         "The following matters must be settled by APRA Legal and the Office of Parliamentary "
         "Counsel before this draft is issued for consultation.",
         "Guidance")
    table(doc, ["Decision", "Required settlement"], S.LEGAL_SETTLEMENT, [4.6, 12.0])

    doc.add_paragraph("Requirements carrying a legal flag", style="Heading 2")
    flagged = [[r["id"], r["title"], r["rationale"]]
               for r in R.REQUIREMENTS if r["legal_flag"]]
    table(doc, ["ID", "Requirement", "Reason for flag"], flagged, [1.5, 4.6, 10.5])

    doc.add_paragraph("Requirements identified as policy design choices", style="Heading 2")
    para(doc,
         "The following requirements are not directly mandated by any comparator authority. They "
         "are proposed by APRA and require consultation and a cost-benefit assessment before issue. "
         "Identification as a policy choice does not indicate that a requirement is unsupported.",
         "Guidance")
    choices = [[r["id"], r["title"], r["rationale"]]
               for r in R.REQUIREMENTS if r["policy_choice"]]
    table(doc, ["ID", "Requirement", "Basis"], choices, [1.5, 4.6, 10.5])


def main():
    doc = Document()
    build_styles(doc)
    page_setup(doc)
    add_header_footer(doc, f"CPS 240 Model Risk Management — {STATUS}",
                      "Australian Prudential Regulation Authority")

    _, req_map, _, _ = G.numbered()

    cover(doc)
    objectives(doc)
    interpretation(doc)
    contents(doc)
    body(doc, req_map)
    annex_a(doc)
    annex_b(doc)
    annex_c(doc)
    annex_d(doc)
    annex_e(doc)
    annex_f(doc)

    doc.save(OUT)
    removed = force_arial(OUT)
    audit = audit_fonts(OUT)

    print(f"wrote {OUT}")
    print(f"  requirements rendered : {audit['requirement_paras']}")
    print(f"  guidance paragraphs   : {audit['guidance_paras']}")
    print(f"  theme font refs removed: {removed}")
    if audit["violations"]:
        print(f"  !! {len(audit['violations'])} font/bold violations")
        for v in audit["violations"][:10]:
            print("    ", v)
    else:
        print("  font/bold audit       : clean")
    return 0 if not audit["violations"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
