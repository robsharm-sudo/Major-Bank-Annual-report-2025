#!/usr/bin/env python3
"""Build CPS XXXX Model Risk Management — the draft cross-industry prudential standard.

Bold paragraphs are mandatory requirements. Everything else is explanatory.
Cross-references into CPG XXXX are computed from the guide's own paragraph
numbering, so the two documents cannot fall out of step.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).parent))

import content_guide as G
import content_requirements as R
import content_sources as S
from docx_common import (
    AMBER, GREY, NAVY, SLATE, audit_fonts, banner, bottom_rule, bullets,
    callout, force_arial, landscape_section, page_break, para, portrait_section,
    page_setup, add_header_footer, build_styles, table, xref_box,
)

OUT = Path(__file__).resolve().parents[1] / "CPS_XXXX_Model_Risk_Management.docx"
AS_OF = "25 July 2026"
STATUS = "POLICY DEVELOPMENT DRAFT — NOT IN FORCE"


def cover(doc):
    p = para(doc, "AUSTRALIAN PRUDENTIAL REGULATION AUTHORITY", "DocMeta")
    p.runs[0].font.size = Pt(9)
    p.runs[0].bold = True

    p = para(doc, "Prudential Standard CPS XXXX", "DocTitle")
    bottom_rule(p)
    para(doc, "Model Risk Management", "DocSubtitle").runs[0].bold = True
    para(doc, "Draft cross-industry prudential standard", "DocMeta")
    para(doc, f"Policy development version — information verified to {AS_OF}", "DocMeta")
    doc.add_paragraph(style="DocMeta")

    banner(doc, STATUS)

    callout(
        doc,
        "Status and legal review",
        "This document is a policy-development draft and is not in force. It has no legal effect. "
        "APRA Legal and the Office of Parliamentary Counsel must settle the enabling provisions, "
        "sectoral application, instrument numbering, definitions, commencement, transition, "
        "enforcement powers and interaction with existing CPS and SPS instruments before "
        "consultation or issue. Annex F records the decisions required. The number 'CPS XXXX' is a "
        "placeholder adopted for drafting convenience and does not indicate an allocated instrument "
        "number.",
    )


def objectives(doc):
    doc.add_paragraph("Objectives and key requirements", style="Heading 1")
    para(doc,
         "The objective of this Prudential Standard is to ensure that an APRA-regulated entity "
         "identifies the models it relies upon, manages the risk that those models are wrong or are "
         "used incorrectly, and can demonstrate that reliance on models is consistent with its "
         "prudential obligations.",
         "Guidance")
    para(doc,
         "Models determine regulatory capital, provisions, valuations, insurance liabilities, unit "
         "prices, credit and pricing decisions, fraud interventions and member outcomes. Model risk "
         "is the risk of adverse consequences from decisions based on incorrect or misused model "
         "outputs. It is a distinct prudential risk and is not adequately managed as a subset of the "
         "risks the models themselves measure.",
         "Guidance")
    para(doc,
         "This Prudential Standard is intended to operate alongside CPS 220 or SPS 220, CPS 230 and "
         "CPS 234, and alongside the model requirements in the capital, actuarial and investment "
         "governance standards. It does not displace those requirements, nor obligations under other "
         "Commonwealth, State or Territory laws or the responsibilities of other regulators.",
         "Guidance")
    para(doc,
         "The key requirements of this Prudential Standard are that an APRA-regulated entity must:",
         "Guidance")
    bullets(doc, [
        "identify the models it uses and maintain an accurate enterprise inventory;",
        "classify models by risk and apply controls of an intensity that matches that classification;",
        "govern models across their full lifecycle, from development through to decommissioning;",
        "subject models to independent validation and to effective challenge;",
        "control how models are used, changed and adjusted;",
        "monitor model performance and act on deterioration;",
        "manage model risk arising from third parties and from artificial intelligence techniques;",
        "assess model risk in aggregate and report it to the Board; and",
        "maintain records and obtain independent assurance sufficient to demonstrate compliance.",
    ])


def interpretation(doc):
    doc.add_paragraph("Interpretation", style="Heading 1")
    callout(
        doc,
        "How to read this Prudential Standard",
        "Every numbered paragraph set in bold is a proposed mandatory requirement. Text that is not "
        "bold is explanatory and does not itself impose an obligation. Each requirement is followed "
        "by explanatory text and by a cross-reference to the paragraphs of CPG XXXX Model Risk "
        "Management that provide guidance on meeting it.",
    )
    para(doc,
         "Proportionality. References in this Prudential Standard to materiality, risk, complexity "
         "and criticality are cumulative considerations, applied to the circumstances of the entity "
         "and of the individual model.",
         "Guidance")
    para(doc,
         "Interaction with other standards. Where a model is also subject to a requirement in a "
         "capital, actuarial or investment governance standard, that requirement continues to apply. "
         "Where a model failure is also an operational risk incident or an information security "
         "incident, the notification requirements of CPS 230 and CPS 234 apply and this Prudential "
         "Standard does not create a separate deadline.",
         "Guidance")
    para(doc,
         "Exercise of APRA powers. Where this Prudential Standard provides for APRA to determine, "
         "approve, require or adjust anything, that power is to be exercised in writing. The "
         "procedural requirements attaching to each such power are matters for legal settlement "
         "and are recorded in Annex F.",
         "Guidance")
    para(doc,
         "Terms defined in Annex A have the meaning given in that Annex.",
         "Guidance")


def front_matter(doc):
    """Authority, application, commencement and the adjustment power.

    Every APRA prudential standard opens with these. Without an adjustments and
    exclusions paragraph in particular there is no lawful route to vary a
    requirement for an individual entity, which would make the instrument far
    more rigid than any comparable standard.
    """
    doc.add_paragraph("Authority, application and commencement", style="Heading 1")

    items = [
        ("A1", "Authority",
         "This Prudential Standard is made under the enabling provisions of the Banking Act 1959, "
         "the Insurance Act 1973, the Life Insurance Act 1995, the Private Health Insurance "
         "(Prudential Supervision) Act 2015 and the Superannuation Industry (Supervision) Act 1993, "
         "as applicable to each class of APRA-regulated entity.",
         ["The precise enabling sections, and whether a single cross-industry instrument can be "
          "made for all classes of APRA-regulated entity or whether companion instruments are "
          "required, must be settled by APRA Legal and the Office of Parliamentary Counsel. "
          "Annex F records this."]),
        ("A2", "Application",
         "This Prudential Standard applies to all APRA-regulated entities, in accordance with the "
         "classes specified in this paragraph and subject to any adjustment or exclusion "
         "determined under paragraph A5.",
         ["The classes to which this Prudential Standard applies, including the treatment of "
          "foreign ADIs and branches, non-operating holding companies, private health insurers "
          "and application at Level 1, Level 2 and Level 3, require settlement before "
          "consultation.",
          "An APRA-regulated entity that is a member of a group must comply with this Prudential "
          "Standard both as an individual entity and, where APRA so determines, on a group basis."]),
        ("A3", "Commencement",
         "This Prudential Standard commences on a date to be determined.",
         ["Commencement and any transitional arrangements are matters for settlement. Given the "
          "uplift required in model identification, inventory completeness and validation "
          "coverage, staged commencement with earlier milestones for inventory and tiering than "
          "for validation would be consistent with the transition periods adopted by comparable "
          "authorities."]),
        ("A4", "Transitional arrangements for models already in use",
         "An APRA-regulated entity must bring a model that is in use at the commencement of this "
         "Prudential Standard into compliance in accordance with a documented remediation plan "
         "approved by the Board or a Board committee, within the period determined by APRA.",
         ["Without this paragraph, requirements expressed as a bar on use — such as the "
          "pre-implementation testing and independent validation requirements — would on "
          "commencement prohibit the continued use of models the entity is already relying on. "
          "That is not the intended effect.",
          "The remediation plan should sequence by model risk tier rather than by convenience, "
          "and should identify the models the entity cannot bring into compliance within the "
          "period and what it proposes to do about them."]),
        ("A5", "Adjustments and exclusions",
         "APRA may, by notice in writing to an APRA-regulated entity, adjust or exclude a specific "
         "requirement in this Prudential Standard in relation to that entity.",
         ["This mirrors the adjustment power in the other cross-industry prudential standards. It "
          "is the route by which a requirement can be varied for an individual entity, and its "
          "absence would make this Prudential Standard less flexible than the instruments it sits "
          "alongside.",
          "The scope of the power, and the procedural requirements attaching to its exercise, "
          "require legal settlement."]),
        ("A6", "Interpretation and previous determinations",
         "Terms defined in Annex A have the meaning given in that Annex, and a reference to a "
         "prudential standard is a reference to that standard as in force from time to time.",
         ["Where an AI system meets the definition of a model, this Prudential Standard applies to "
          "it. The interaction with any separate prudential standard on artificial intelligence "
          "risk management is addressed in M34 and is a matter for settlement once the numbering "
          "and scope of that instrument are known."]),
    ]

    for pid, title, operative, notes in items:
        doc.add_paragraph(f"{pid} — {title}", style="Heading 3")
        p = doc.add_paragraph(style="Requirement")
        p.add_run(f"{pid}. {operative}").bold = True
        for n in notes:
            para(doc, n, "Guidance")

    callout(
        doc,
        "Front matter is drafting scaffolding",
        "Paragraphs A1 to A6 set out the instrument mechanics that every APRA prudential standard "
        "carries. They are drafted here so that the standard is structurally complete and can be "
        "read as an instrument, but the enabling provisions, application classes, commencement, "
        "transition period and the scope of the adjustment power are all matters for APRA Legal "
        "and the Office of Parliamentary Counsel. Annex F records each of them.",
    )


def contents(doc):
    doc.add_paragraph("Contents", style="Heading 1")
    para(doc, "Authority, application and commencement  (A1–A6)", "ContentsHead")
    seen = []
    for r in R.REQUIREMENTS:
        if r["part"] not in seen:
            seen.append(r["part"])
    for part in seen:
        para(doc, part, "ContentsHead")
        secs = []
        for r in R.REQUIREMENTS:
            if r["part"] == part and r["section"] not in secs:
                secs.append(r["section"])
        for s in secs:
            ids = [r["id"] for r in R.REQUIREMENTS if r["section"] == s]
            para(doc, f"{s}  ({ids[0]}–{ids[-1]})" if len(ids) > 1
                 else f"{s}  ({ids[0]})", "ContentsItem")
    doc.add_paragraph(style="DocMeta")
    for a in ["Annex A — Definitions",
              "Annex B — Minimum model inventory fields",
              "Annex C — Model risk tiering indicators",
              "Annex D — International comparison of model risk requirements",
              "Annex E — Source register",
              "Annex F — Matters requiring legal and policy settlement"]:
        para(doc, a, "ContentsHead")


def body(doc, req_map):
    current_part = None
    current_section = None

    first = True
    for r in R.REQUIREMENTS:
        if r["part"] != current_part:
            # Only the first operative part starts a fresh page; the rest flow, so
            # the standard reads continuously rather than as ten short chapters.
            if first:
                page_break(doc)
                first = False
            doc.add_paragraph(r["part"], style="Heading 1")
            current_part = r["part"]
            current_section = None

        if r["section"] != current_section:
            doc.add_paragraph(r["section"], style="Heading 2")
            current_section = r["section"]

        doc.add_paragraph(f"{r['id']} — {r['title']}", style="Heading 3")

        p = doc.add_paragraph(style="Requirement")
        run = p.add_run(f"{r['id']}. {r['requirement']}")
        run.bold = True

        for g in r["guidance"]:
            para(doc, g, "Guidance")

        nums = req_map.get(r["id"], [])
        if nums:
            xr = para(doc, f"→ Guidance: CPG XXXX paragraphs {G.ranges(nums)}.", "CrossRef")
            xr.runs[0].italic = True


def annex_a(doc):
    page_break(doc)
    doc.add_paragraph("Annex A — Definitions", style="Heading 1")
    para(doc,
         "This Annex is intended to form part of the final Prudential Standard. Definitions are "
         "subject to legal drafting and to alignment with Australian Government and APRA "
         "terminology.",
         "Guidance")
    table(doc, ["Term", "Draft definition"], S.DEFINITIONS, [4.0, 12.6])


def annex_b(doc):
    page_break(doc)
    doc.add_paragraph("Annex B — Minimum model inventory fields", style="Heading 1")
    para(doc,
         "M13 requires the enterprise model inventory to record at least the following for each "
         "model. These are minimum fields; an entity should record whatever further information its "
         "framework requires.",
         "Guidance")
    table(doc, ["Field group", "Minimum content"], S.INVENTORY_FIELDS, [4.0, 12.6])


def annex_c(doc):
    page_break(doc)
    doc.add_paragraph("Annex C — Model risk tiering indicators", style="Heading 1")
    para(doc,
         "An entity would ordinarily assign a model to its highest risk tier where one or more of "
         "the following indicators is material. Classification remains outcome-based and "
         "context-specific, and no single indicator is determinative.",
         "Guidance")
    for b in S.TIER_INDICATORS:
        para(doc, "• " + b, "Bullet")


def annex_d(doc):
    """The wide international comparison. Landscape so the columns are readable."""
    landscape_section(doc)
    doc.add_paragraph("Annex D — International comparison of model risk requirements",
                      style="Heading 1")
    para(doc,
         "This Annex compares the model risk requirements and expectations of the authorities whose "
         "frameworks informed this Prudential Standard. Legal status differs materially between "
         "them and is stated in each cell group. An overseas instrument is not binding in Australia "
         "whatever its status in its home jurisdiction.",
         "Guidance")
    para(doc,
         "Status codes: [B] binding in the home jurisdiction · [SE] supervisory expectation, not "
         "binding law · [G] guidance · [PB] proposed or consultative · [A] analytical or advisory · "
         "[—] not addressed.",
         "SourceNote")

    headers = ["Domain", "APRA (proposed CPS XXXX)", "OCC / Federal Reserve (US)",
               "PRA (UK)", "OSFI (Canada)", "ECB / SSM (EU)", "MAS (Singapore)",
               "BCBS", "FSB"]
    widths = [2.5, 3.1, 3.0, 3.0, 3.0, 3.0, 2.9, 2.7, 2.5]
    table(doc, headers, S.CROSSWALK_ROWS, widths, font_pt=7)

    portrait_section(doc)


def annex_e(doc):
    doc.add_paragraph("Annex E — Source register", style="Heading 1")
    para(doc,
         "Every source relied upon in this Prudential Standard and in CPG XXXX. URLs are to official "
         "authority websites and were checked on the date shown. Where a source could not be "
         "confirmed from an official domain it is marked accordingly and has not been relied upon "
         "for any statement of requirement.",
         "Guidance")
    rows = [[s["id"], s["authority"], s["title"], s["date"], s["status"], s["url"]]
            for s in S.SOURCES]
    table(doc, ["ID", "Authority", "Title", "Date / effect", "Legal status", "Official URL"],
          rows, [1.9, 1.5, 4.4, 2.3, 3.0, 3.5], font_pt=7)


def annex_f(doc):
    page_break(doc)
    doc.add_paragraph("Annex F — Matters requiring legal and policy settlement", style="Heading 1")
    para(doc,
         "The following matters must be settled by APRA Legal and the Office of Parliamentary "
         "Counsel before this draft is issued for consultation.",
         "Guidance")
    table(doc, ["Decision", "Required settlement"], S.LEGAL_SETTLEMENT, [4.6, 12.0])

    doc.add_paragraph("Requirements carrying a legal flag", style="Heading 2")
    flagged = [[r["id"], r["title"], r["rationale"]]
               for r in R.REQUIREMENTS if r["legal_flag"]]
    table(doc, ["ID", "Requirement", "Reason for flag"], flagged, [1.5, 4.6, 10.5])

    doc.add_paragraph("Requirements identified as policy design choices", style="Heading 2")
    para(doc,
         "The following requirements are not directly mandated by any comparator authority. They "
         "are proposed by APRA and require consultation and a cost-benefit assessment before issue. "
         "Identification as a policy choice does not indicate that a requirement is unsupported.",
         "Guidance")
    choices = [[r["id"], r["title"], r["rationale"]]
               for r in R.REQUIREMENTS if r["policy_choice"]]
    table(doc, ["ID", "Requirement", "Basis"], choices, [1.5, 4.6, 10.5])


def main():
    doc = Document()
    build_styles(doc)
    page_setup(doc)
    add_header_footer(doc, f"CPS XXXX Model Risk Management — {STATUS}",
                      "Australian Prudential Regulation Authority")

    _, req_map, _, _ = G.numbered()

    cover(doc)
    objectives(doc)
    interpretation(doc)
    front_matter(doc)
    contents(doc)
    body(doc, req_map)
    annex_a(doc)
    annex_b(doc)
    annex_c(doc)
    annex_d(doc)
    annex_e(doc)
    annex_f(doc)

    doc.save(OUT)
    removed = force_arial(OUT)
    audit = audit_fonts(OUT)

    print(f"wrote {OUT}")
    print(f"  requirements rendered : {audit['requirement_paras']}")
    print(f"  guidance paragraphs   : {audit['guidance_paras']}")
    print(f"  theme font refs removed: {removed}")
    if audit["violations"]:
        print(f"  !! {len(audit['violations'])} font/bold violations")
        for v in audit["violations"][:10]:
            print("    ", v)
    else:
        print("  font/bold audit       : clean")
    return 0 if not audit["violations"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
