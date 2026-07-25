"""Shared workbook furniture for the model risk management assessment workbook.

Every sheet is built the same way: a title row, a status strapline, a blank row,
then a frozen and filtered header row at row 4 with data from row 5. That layout
is inherited deliberately from the AI package workbook supplied as input, so the
two workbooks can sit side by side without a reader having to relearn anything.
"""

from __future__ import annotations

from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ARIAL = "Arial"

NAVY = "1B335F"
SLATE = "2F5497"
LIGHT = "EEF2F8"
ALT = "F7F9FC"
AMBER = "9C6300"
AMBER_FILL = "FFF4DC"
RED_FILL = "FCE4E4"
GREEN_FILL = "E4F3E7"
GREY = "595959"
INPUT_BLUE = "0000FF"

HEADER_ROW = 4
FIRST_DATA_ROW = 5

thin = Side(style="thin", color="C7D0DF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def sheet_frame(ws, title: str, strapline: str, headers, widths, *,
                freeze_col: str = "A", tab_colour: str = NAVY,
                wrap_from: int = 1):
    """Lay down the standard title / strapline / header frame and return the sheet."""
    ws.sheet_properties.tabColor = tab_colour

    ws["A1"] = title
    ws["A1"].font = Font(name=ARIAL, size=14, bold=True, color=NAVY)

    ws["A2"] = strapline
    ws["A2"].font = Font(name=ARIAL, size=9, italic=True, color=AMBER)

    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=h)
        c.font = Font(name=ARIAL, size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = BORDER

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[HEADER_ROW].height = 30
    ws.freeze_panes = f"{freeze_col}{FIRST_DATA_ROW}"
    ws.auto_filter.ref = (
        f"A{HEADER_ROW}:{get_column_letter(len(headers))}{HEADER_ROW}"
    )
    return ws


def write_rows(ws, rows, *, start=FIRST_DATA_ROW, wrap_cols=None, zebra=True,
               size=9, row_height=None):
    """Write data rows with consistent typography and optional wrapping."""
    wrap_cols = set(wrap_cols or [])
    for r, row in enumerate(rows, start=start):
        for i, val in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=val)
            c.font = Font(name=ARIAL, size=size)
            c.border = BORDER
            c.alignment = Alignment(
                vertical="top",
                wrap_text=(i in wrap_cols) if wrap_cols else True,
            )
            if zebra and (r - start) % 2 == 1:
                c.fill = PatternFill("solid", fgColor=ALT)
        if row_height:
            ws.row_dimensions[r].height = row_height
    return start + len(rows) - 1


def note(ws, cell: str, text: str, author: str = "Model risk package") -> None:
    """Attach an explanatory comment — used to document every scoring judgement."""
    c = ws[cell]
    cm = Comment(text, author)
    cm.width = 320
    cm.height = 160
    c.comment = cm


def label(ws, cell: str, text: str, *, bold=False, size=9, colour="000000",
          italic=False, fill=None, wrap=True):
    c = ws[cell]
    c.value = text
    c.font = Font(name=ARIAL, size=size, bold=bold, color=colour, italic=italic)
    c.alignment = Alignment(vertical="top", wrap_text=wrap)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    return c


def input_cell(ws, cell: str, value):
    """Blue text + yellow fill: the convention for a cell the reader may change."""
    c = ws[cell]
    c.value = value
    c.font = Font(name=ARIAL, size=9, bold=True, color=INPUT_BLUE)
    c.fill = PatternFill("solid", fgColor="FFFF99")
    c.border = BORDER
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c


def score_scale(ws, ref: str) -> None:
    """Red-amber-green scale for 0-5 capability scores."""
    ws.conditional_formatting.add(
        ref,
        ColorScaleRule(
            start_type="num", start_value=0, start_color="F4B7B7",
            mid_type="num", mid_value=2.5, mid_color="FFE9A8",
            end_type="num", end_value=5, end_color="B7DDBE",
        ),
    )


def gap_scale(ws, ref: str) -> None:
    """Inverted scale: a large gap is the thing to worry about."""
    ws.conditional_formatting.add(
        ref,
        ColorScaleRule(
            start_type="num", start_value=0, start_color="B7DDBE",
            mid_type="num", mid_value=1.5, mid_color="FFE9A8",
            end_type="num", end_value=3.5, end_color="F4B7B7",
        ),
    )


def priority_bands(ws, ref: str) -> None:
    for text, fill, font_colour in (
        ("Critical", "F4B7B7", "8B0000"),
        ("High", "FBD9B5", "8A4B00"),
        ("Medium", "FFF0B8", "6B5300"),
        ("Low", "D8ECDC", "1E5B2A"),
    ):
        ws.conditional_formatting.add(
            ref,
            CellIsRule(
                operator="equal",
                formula=[f'"{text}"'],
                fill=PatternFill("solid", fgColor=fill),
                font=Font(name=ARIAL, size=9, bold=True, color=font_colour),
            ),
        )


def status_bands(ws, ref: str) -> None:
    mapping = (
        ("REFUTED", "F4B7B7", "8B0000"),
        ("PARTIALLY_WRONG", "FBD9B5", "8A4B00"),
        ("UNSUPPORTED", "FFF0B8", "6B5300"),
        ("CONFIRMED", "D8ECDC", "1E5B2A"),
        ("High", "F4B7B7", "8B0000"),
        ("Medium", "FFF0B8", "6B5300"),
        ("Low", "D8ECDC", "1E5B2A"),
    )
    for text, fill, font_colour in mapping:
        ws.conditional_formatting.add(
            ref,
            CellIsRule(
                operator="equal",
                formula=[f'"{text}"'],
                fill=PatternFill("solid", fgColor=fill),
                font=Font(name=ARIAL, size=9, bold=True, color=font_colour),
            ),
        )


def verification_bands(ws, ref: str) -> None:
    for text, fill, font_colour in (
        ("Verified — primary source", "D8ECDC", "1E5B2A"),
        ("Verified — secondary source", "FFF0B8", "6B5300"),
        ("Unverified — not independently confirmed", "F4B7B7", "8B0000"),
    ):
        ws.conditional_formatting.add(
            ref,
            CellIsRule(
                operator="equal",
                formula=[f'"{text}"'],
                fill=PatternFill("solid", fgColor=fill),
                font=Font(name=ARIAL, size=9, bold=True, color=font_colour),
            ),
        )


def section_head(ws, row: int, text: str, span: int = 8) -> int:
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=SLATE)
    c.alignment = Alignment(vertical="center")
    for i in range(2, span + 1):
        ws.cell(row=row, column=i).fill = PatternFill("solid", fgColor=SLATE)
    ws.row_dimensions[row].height = 18
    return row + 1
