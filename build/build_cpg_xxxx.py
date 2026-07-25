#!/usr/bin/env python3
"""Build CPG XXXX Model Risk Management — the draft prudential practice guide.

Follows the construction of APG 250, supplied as a model: numbered "should"
paragraphs, a cross-reference strip opening each chapter, source attribution
lines carrying legal-status codes, and amber verification annotations.
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(Path(__file__).parent))

import content_guide as G
import content_requirements as R
import content_sources as S
from docx_common import (
    audit_fonts, banner, bottom_rule, bullets, callout, force_arial,
    landscape_section, page_break, para, portrait_section, page_setup,
    add_header_footer, build_styles, table, xref_box,
)

OUT = Path(__file__).resolve().parents[1] / "CPG_XXXX_Model_Risk_Management.docx"
AS_OF = "25 July 2026"
STATUS = "POLICY DEVELOPMENT DRAFT — NOT IN FORCE"

REQ_BY_ID = {r["id"]: r for r in R.REQUIREMENTS}


def cover(doc):
    p = para(doc, "AUSTRALIAN PRUDENTIAL REGULATION AUTHORITY", "DocMeta")
    p.runs[0].font.size = Pt(9)
    p.runs[0].bold = True

    p = para(doc, "Prudential Practice Guide CPG XXXX", "DocTitle")
    bottom_rule(p)
    para(doc, "Model Risk Management", "DocSubtitle").runs[0].bold = True
    para(doc, "Draft cross-industry prudential practice guide", "DocMeta")
    para(doc, f"Policy development version — information verified to {AS_OF}", "DocMeta")
    doc.add_paragraph(style="DocMeta")

    banner(doc, STATUS)

    callout(
        doc,
        "How to read this guide",
        "This guide uses 'should' to indicate APRA's expectations. Compliance with a practice guide "
        "is not mandatory, but APRA expects an entity either to apply the guidance or to be able to "
        "explain the alternative approach it has adopted and why that approach meets the "
        "requirements of CPS XXXX. Requirements use 'must' and appear only in CPS XXXX. Verification "
        "annotations appear in amber throughout, recording the evidential basis for statements "
        "about overseas frameworks.",
    )


def chapter_xref(doc, chapter_name, chapter_reqs):
    ids = sorted(chapter_reqs.get(chapter_name, []))
    if not ids:
        return
    titles = ", ".join(ids)
    xref_box(doc, f"CPS XXXX cross-reference: {titles}")


def render_items(doc, flat, chapter_reqs):
    for it in flat:
        k = it["k"]

        if k == "chapter":
            page_break(doc)
            doc.add_paragraph(it["t"], style="Heading 1")
            chapter_xref(doc, it["t"], chapter_reqs)
            if it.get("intro"):
                para(doc, it["intro"], "GuideBody")

        elif k == "h":
            doc.add_paragraph(it["t"], style="Heading 2")

        elif k == "p":
            p = doc.add_paragraph(style="GuideBody")
            n = p.add_run(f"{it['n']}. ")
            n.bold = True
            p.add_run(it["t"])

        elif k == "bul":
            for b in it["items"]:
                para(doc, "• " + b, "Bullet")

        elif k == "src":
            para(doc, it["t"], "SourceNote")

        elif k == "ver":
            para(doc, "⚠ " + it["t"], "Verification")

        elif k == "tbl":
            table(doc, it["headers"], it["rows"], it["widths"], font_pt=8)
            if it.get("caption"):
                para(doc, it["caption"], "Caption")


def annex_traceability(doc, req_map):
    page_break(doc)
    doc.add_paragraph("Annex A — Requirement to guidance traceability", style="Heading 1")
    para(doc,
         "Every mandatory requirement in CPS XXXX and the paragraphs of this guide that address it. "
         "Requirements with no guidance paragraph would indicate a gap in this guide; there are none.",
         "GuideBody")
    rows = []
    for r in R.REQUIREMENTS:
        nums = req_map.get(r["id"], [])
        rows.append([r["id"], r["title"], G.ranges(nums) or "—",
                     r["provenance"],
                     "Yes" if r["policy_choice"] else "No"])
    table(doc, ["CPS XXXX", "Requirement", "CPG XXXX paragraphs", "Provenance", "Policy choice"],
          rows, [1.5, 6.4, 3.4, 3.3, 2.0], font_pt=8)


def annex_principles(doc):
    landscape_section(doc)
    doc.add_paragraph("Annex B — Mapping to international principles", style="Heading 1")
    para(doc,
         "Where each CPS XXXX requirement sits against the enumerated principles, outcomes and core "
         "elements of the comparator frameworks. Principle labels are reproduced as printed by the "
         "issuing authority. A blank cell indicates that the framework does not address the "
         "requirement through an enumerated principle, not that it is silent on the subject.",
         "GuideBody")
    rows = [[r["id"], r["title"], r["principles"]] for r in R.REQUIREMENTS]
    table(doc, ["CPS XXXX", "Requirement", "Corresponding international principle"],
          rows, [1.8, 6.0, 16.2], font_pt=7.5)
    portrait_section(doc)


def annex_frameworks(doc):
    doc.add_paragraph("Annex C — The comparator frameworks in outline", style="Heading 1")
    para(doc,
         "The enumerated structure of each framework relied upon in this guide, as printed by the "
         "issuing authority. This Annex is provided so that the attribution lines throughout the "
         "guide can be traced to a specific principle.",
         "GuideBody")
    for fw in S.PRINCIPLE_FRAMEWORKS:
        doc.add_paragraph(fw["instrument"], style="Heading 2")
        para(doc, fw["note"], "SourceNote")
        table(doc, ["Label", "Title as printed"], fw["items"], [3.4, 13.2], font_pt=8)


def annex_status(doc):
    page_break(doc)
    doc.add_paragraph("Annex D — Legal status of sources", style="Heading 1")
    para(doc,
         "The legal character of each source relied upon. This matters because an instrument that "
         "is binding in its home jurisdiction is still not binding in Australia, and several of the "
         "most influential model risk instruments are supervisory guidance rather than law even at "
         "home.",
         "GuideBody")
    rows = [[s["id"], s["authority"], s["title"], s["status"], s["scope"]] for s in S.SOURCES]
    table(doc, ["ID", "Authority", "Title", "Legal status", "Scope"],
          rows, [1.7, 1.4, 4.6, 4.4, 4.5], font_pt=7)


def main():
    doc = Document()
    build_styles(doc)
    page_setup(doc)
    add_header_footer(doc, f"CPG XXXX Model Risk Management — {STATUS}",
                      "Australian Prudential Regulation Authority")

    flat, req_map, chapter_reqs, total = G.numbered()

    cover(doc)
    render_items(doc, flat, chapter_reqs)
    annex_traceability(doc, req_map)
    annex_principles(doc)
    annex_frameworks(doc)
    annex_status(doc)

    doc.save(OUT)
    removed = force_arial(OUT)
    audit = audit_fonts(OUT)

    print(f"wrote {OUT}")
    print(f"  numbered guidance paragraphs: {total}")
    print(f"  chapters                    : {len(G.GUIDE)}")
    print(f"  theme font refs removed     : {removed}")
    if audit["violations"]:
        print(f"  !! {len(audit['violations'])} font violations")
        for v in audit["violations"][:10]:
            print("    ", v)
    else:
        print("  font audit                  : clean")
    return 0 if not audit["violations"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
